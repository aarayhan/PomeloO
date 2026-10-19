"""End-to-end check: every catalog conversion pair and every tool, on generated fixtures.

Run:  .venv\\Scripts\\python.exe tests\\test_convert.py   (prints a table, exits 1 on failure)
Slides (ppt/pptx) need LibreOffice or PowerPoint; they are reported as SKIP when neither exists.
"""
import shutil
import sys
import tempfile
import traceback
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pymupdf as fitz  # noqa: E402
from PIL import Image, ImageDraw  # noqa: E402

from pomeloo.catalog import all_input_exts, targets, tools  # noqa: E402
from pomeloo.engines import convert  # noqa: E402
from pomeloo.engines.media import ff  # noqa: E402
import pomeloo.engines.image  # noqa: E402,F401  (registers HEIF)
from pomeloo.tools import run_tool  # noqa: E402


def make_fixtures(d: Path) -> dict:
    f = {}
    img = Image.new("RGB", (320, 240), "#ff8a00")
    ImageDraw.Draw(img).ellipse((60, 40, 260, 200), fill="#2050ff")
    for e, fmt in [("jpg", "JPEG"), ("png", "PNG"), ("webp", "WEBP"), ("heic", "HEIF"),
                   ("tiff", "TIFF"), ("avif", "AVIF"), ("bmp", "BMP")]:
        f[e] = d / f"sample.{e}"
        img.save(f[e], fmt)
    f["svg"] = d / "sample.svg"
    f["svg"].write_text('<svg xmlns="http://www.w3.org/2000/svg" width="100" height="80">'
                        '<rect width="100" height="80" fill="orange"/></svg>')
    f["wav"] = d / "sample.wav"
    ff("-f", "lavfi", "-i", "sine=frequency=440:duration=3", f["wav"])
    for e in ["mp3", "m4a", "flac", "ogg", "opus", "aiff", "wma"]:
        f[e] = convert(f["wav"], e)
    f["mp4"] = d / "sample.mp4"
    ff("-f", "lavfi", "-i", "testsrc=size=320x240:rate=25:duration=3", "-f", "lavfi", "-i",
       "sine=duration=3", "-shortest", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", f["mp4"])
    for e in ["mov", "mkv", "webm", "avi", "wmv", "gif"]:
        f[e] = convert(f["mp4"], e)
    doc = fitz.open()
    for i in range(2):
        page = doc.new_page()
        page.insert_text((72, 72), f"Hello page {i + 1}", fontsize=24)
        page.insert_image(fitz.Rect(72, 120, 300, 300), filename=str(f["png"]))
    f["pdf"] = d / "sample.pdf"
    doc.save(f["pdf"])
    f["md"] = d / "sample.md"
    f["md"].write_text("# Title\n\nSome *rich* text.\n\n| a | b |\n|---|---|\n| 1 | 2 |\n")
    f["txt"] = d / "sample.txt"
    f["txt"].write_text("Plain text.\n\nSecond paragraph.")
    f["html"] = d / "sample.html"
    f["html"].write_text("<h1>Hi</h1><p>para</p>")
    f["docx"], f["epub"] = convert(f["md"], "docx"), convert(f["md"], "epub")
    f["rtf"] = d / "sample.rtf"
    f["rtf"].write_text(r"{\rtf1\ansi Hello {\b bold} world.}")
    f["rtfd"] = d / "sample.rtfd"
    f["rtfd"].mkdir()
    shutil.copy(f["rtf"], f["rtfd"] / "TXT.rtf")
    f["pages"] = d / "sample.pages"
    with zipfile.ZipFile(f["pages"], "w") as z:
        z.write(f["pdf"], "QuickLook/Preview.pdf")
    f["csv"] = d / "sample.csv"
    f["csv"].write_text("name,qty\napple,3\npear,5\n")
    f["tsv"], f["xlsx"], f["ods"] = (convert(f["csv"], e) for e in ("tsv", "xlsx", "ods"))
    f["srt"] = d / "sample.srt"
    f["srt"].write_text("1\n00:00:01,000 --> 00:00:02,500\nHello <i>there</i>\n\n2\n00:00:03,000 --> 00:00:04,000\nBye\n")
    f["vtt"] = convert(f["srt"], "vtt")
    f["zip"] = d / "sample.zip"
    with zipfile.ZipFile(f["zip"], "w") as z:
        z.write(f["txt"], "folder/a.txt")
    f["tar"], f["gz"] = convert(f["zip"], "tar"), convert(f["zip"], "gz")
    try:
        from pptx import Presentation
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = "Slide"
        f["pptx"] = d / "sample.pptx"
        prs.save(f["pptx"])
    except Exception:
        pass
    # Formats we can read but not write: real files kept in tests/fixtures
    # (xls/ppt saved by WPS Office as 97-2003, rar from libarchive's test suite).
    for e in ("xls", "ppt", "rar"):
        src = Path(__file__).parent / "fixtures" / f"sample.{e}"
        if src.exists():
            f[e] = d / src.name
            shutil.copy(src, f[e])
    return f


WORDS = ("title", "rich", "plain", "hello", "page", "hi", "para", "bold", "apple", "name", "slide")


def _text(p: Path, e: str) -> str:
    """Readable text of a document output, extracted the way a real reader would."""
    if e == "pdf":
        return " ".join(pg.get_text() for pg in fitz.open(p))
    if e in ("docx", "epub", "pptx", "ods"):
        with zipfile.ZipFile(p) as z:
            return " ".join(z.read(n).decode("utf-8", "replace") for n in z.namelist()
                            if n.endswith((".xml", ".xhtml", ".html")))
    if e == "xlsx":
        from openpyxl import load_workbook
        return " ".join(str(c) for ws in load_workbook(p).worksheets for row in ws.iter_rows(values_only=True)
                        for c in row if c is not None)
    return p.read_text(encoding="utf-8", errors="strict")


def validate(p):
    """Open the output with a real reader for its type; raise AssertionError if it is not usable."""
    import csv
    import subprocess
    import tarfile
    from pomeloo.catalog import ext_of, kind
    from pomeloo.engines.image import open_image
    from pomeloo.engines.media import FF, probe
    p = Path(p)
    assert p.exists(), "missing"
    if p.is_dir():
        files = [f for f in p.rglob("*") if f.is_file()]
        assert files, "empty folder"
        for f in files:
            if ext_of(f) in ("png", "jpg", "pdf", "mp4", "mov", "mkv", "webm", "avi", "wmv"):
                validate(f)
        return True
    e = ext_of(p)
    assert p.stat().st_size > 0, "empty file"
    if e in ("jpg", "png", "webp", "heic", "tiff", "avif", "bmp"):
        im = open_image(p)
        assert im.width > 0 and im.height > 0
    elif kind(e) in ("audio", "video"):
        r = subprocess.run([FF, "-v", "error", "-i", str(p), "-f", "null", "-"], capture_output=True, text=True)
        assert r.returncode == 0 and not r.stderr.strip(), f"decode error: {r.stderr.strip()[:120]}"
        assert probe(p)["duration"] > 0.2, "no duration"
    elif e in ("zip",):
        with zipfile.ZipFile(p) as z:
            assert z.testzip() is None and z.namelist()
    elif e in ("tar", "gz"):
        with tarfile.open(p) as t:
            assert t.getmembers()
    elif e in ("srt", "vtt"):
        from pomeloo.engines.misc import read_cues
        assert read_cues(p), "no cues"
    elif e in ("csv", "tsv"):
        rows = list(csv.reader(p.open(encoding="utf-8"), delimiter="\t" if e == "tsv" else ","))
        assert rows and any(w in " ".join(sum(rows, [])).lower() for w in WORDS), "content lost"
    elif e == "epub":
        with zipfile.ZipFile(p) as z:
            assert z.read("mimetype") == b"application/epub+zip"
    elif e == "pptx":
        from pptx import Presentation
        assert len(Presentation(p).slides) >= 1
    elif e == "ods":
        import pandas as pd
        assert len(pd.read_excel(p, engine="odf", sheet_name=None))
    if e in ("docx", "epub", "xlsx", "ods", "txt", "md", "html"):  # text survived the conversion
        assert any(w in _text(p, e).lower() for w in WORDS), "content lost"
    if e == "pdf":  # image-only PDFs (from photos) have no text, so only check that it renders
        doc = fitz.open(p)
        assert doc.page_count >= 1
        doc[0].get_pixmap(dpi=20)  # renders
    return True


def ok(p) -> bool:
    try:
        return validate(p)
    except AssertionError as e:
        print(f"  invalid {Path(p).name}: {e}")
        return False


TOOL_PARAMS = {
    "annotateImage": {"shapes": [{"kind": k, "points": [[0.2, 0.2], [0.6, 0.7]], "text": "Hi", "n": 1}
                                 for k in ("pen", "highlight", "arrow", "line", "rect", "ellipse", "number",
                                           "text", "magnify")]},
    "cropImage": {"rect": [0.1, 0.1, 0.5, 0.5]},
    "redactImage": {"rects": [[0.2, 0.2, 0.3, 0.3]]},
    "redactVideo": {"rects": [[0.2, 0.2, 0.3, 0.3]]},
    "cropVideo": {"rect": [0, 0, 0.5, 0.5]},
    "trimAudio": {"start": 0.5, "end": 2.0},
    "trimVideo": {"start": 0.5, "end": 2.0},
    "organizePDF": {"order": [1, 0, 0], "rotations": {"2": 90}},  # reorder + duplicate + rotate
    "removeMetadata": {"mode": "remove"},
}


# Extra option combinations for the tool windows (docs/features.md "Tool details").
EXTRA = [
    ("png", "compress", {"preset": "smallest", "resize": 50}),
    ("mp4", "compress", {"percent": 60, "resize": 50}),
    ("png", "editImage", {"exposure": 0.5, "contrast": 20, "highlights": -30, "shadows": 30, "saturation": 20,
                          "temperature": 30, "tint": 10, "clarity": 40, "dehaze": 30, "sharpness": 30,
                          "noise_reduction": 30, "grain": 20, "vignette": 40, "rotate": 90, "flip": "horizontal"}),
    ("png", "frameImage", {"background": "gradient", "aspect": "4:5", "radius": 10, "shadow": 50}),
    ("png", "frameImage", {"background": "blur", "aspect": "16:9"}),
    ("png", "frameImage", {"background": "image", "image": "@jpg"}),
    ("png", "cropImage", {"width": 100, "height": 100}),
    ("png", "redactImage", {"auto": "both", "style": "color", "color": "#ff0000"}),
    ("png", "createCollage*", {"layout": "featured", "radius": 20, "width": 800}),
    ("png", "createCollage*", {"layout": "row"}),
    ("png", "createCollage*", {"layout": "column"}),
    ("wav", "normalizeAudio", {"loudness": -20, "lra": 7, "true_peak": -2}),
    ("wav", "audioToVideo", {"style": "image", "image": "@jpg", "shape": "portrait"}),
    ("wav", "audioToVideo", {"style": "spectrum", "shape": "square"}),
    ("wav", "audioChannels", {"channels": 2, "left": 50, "right": 150}),
    ("wav", "trimAudio", {"start": 0.2, "end": 2.5, "trim_silence": True}),
    ("mp4", "videoSnapshots", {"times": [0.5, 1.0, 2.0]}),
    ("mp4", "splitVideo", {"points": [1.0, 2.2], "precise": True}),
    ("mp4", "splitVideo", {"parts": 3, "precise": False}),
    ("mp4", "redactVideo", {"rects": [[0.1, 0.1, 0.2, 0.2], [0.5, 0.5, 0.2, 0.2]], "style": "color",
                            "times": [[0.5, 1.5], None]}),
    ("mp4", "cropVideo", {"width": 200, "height": 200}),
    ("mp4", "joinVideos*", {}),
]


def main():
    results = []
    with tempfile.TemporaryDirectory() as t:
        d = Path(t)
        fx = make_fixtures(d)
        for ext in all_input_exts():
            if ext not in fx:
                results.append(("SKIP", f"{ext} (no fixture: needs a real {ext} file)"))
                continue
            for tgt in targets(ext):
                try:
                    out = convert(fx[ext], tgt)
                    results.append(("OK" if ok(out) else "FAIL", f"{ext} -> {tgt}"))
                except Exception as e:
                    status = "SKIP" if "LibreOffice" in str(e) else "FAIL"
                    results.append((status, f"{ext} -> {tgt}: {e}"))
            for n in (1, 2):
                for tid in tools(ext, n):
                    if n == 2 and tid in tools(ext, 1) and tid not in ("mergePDF", "joinVideos"):
                        continue  # per-file tools already covered by n=1
                    paths = [fx[ext]] * n
                    try:
                        outs = run_tool(tid, paths, TOOL_PARAMS.get(tid))
                        good = all(isinstance(o, str) or ok(o) for o in outs)
                        results.append(("OK" if good else "FAIL", f"{ext} tool {tid} (n={n})"))
                    except Exception as e:
                        traceback.print_exc()
                        results.append(("FAIL", f"{ext} tool {tid} (n={n}): {e}"))
        silent = d / "silent.mp4"
        ff("-f", "lavfi", "-i", "testsrc=size=160x120:duration=1", "-pix_fmt", "yuv420p", silent)
        for ext, tid, params in EXTRA:
            multi = tid.endswith("*")
            tid = tid.rstrip("*")
            params = {k: str(fx[v[1:]]) if isinstance(v, str) and v.startswith("@") else v for k, v in params.items()}
            paths = [fx[ext], silent] if tid == "joinVideos" else [fx[ext]] * (3 if multi else 1)
            try:
                outs = run_tool(tid, paths, params)
                good = all(isinstance(o, str) or ok(o) for o in outs)
                results.append(("OK" if good else "FAIL", f"extra {tid} {params}"))
            except Exception as e:
                traceback.print_exc()
                results.append(("FAIL", f"extra {tid} {params}: {e}"))
    for status, msg in results:
        if status != "OK":
            print(f"{status:5} {msg}")
    counts = {s: sum(1 for r in results if r[0] == s) for s in ("OK", "FAIL", "SKIP")}
    print(counts)
    return 1 if counts["FAIL"] else 0


if __name__ == "__main__":
    sys.exit(main())
