"""Exhibition mode (tray -> Exhibition mode): turns one laptop into a demo booth.

- KeyCaster: the modifier keys and shortcuts being pressed, big, above the taskbar, so onlookers see
  "Shift" or "Win + Ctrl + T". Plain typing (and Shift+letter) is never shown.
- Attract card: after a while without input, three things to try and a QR code to the website.
- Demo folder: Desktop\\PomeloO Demo with sample files, rebuilt between visitors."""
import ctypes
import os
import threading
from pathlib import Path

from PySide6.QtCore import QLocale, QObject, QStandardPaths, Qt, QTimer, Signal
from PySide6.QtGui import QCursor, QGuiApplication, QImage, QPixmap
from PySide6.QtWidgets import QFrame, QHBoxLayout, QLabel, QVBoxLayout, QWidget

SITE = "https://pomeloo.netlify.app"
IDLE_S = 25        # seconds without input before the attract card shows
MARK = ".pomeloo-demo"  # only a folder holding this file is ever emptied
ORDER = ["Win", "Ctrl", "Alt", "Shift"]
VKS = ([(0x5B, "Win"), (0x5C, "Win"), (0x11, "Ctrl"), (0x12, "Alt"), (0x10, "Shift"), (0x20, "Space"),
        (0x09, "Tab"), (0x0D, "Enter"), (0x1B, "Esc")] + [(c, chr(c)) for c in range(0x30, 0x5B) if not 0x39 < c < 0x41]
       + [(0x70 + i, f"F{i + 1}") for i in range(12)])
FLAGS = Qt.FramelessWindowHint | Qt.WindowStaysOnTopHint | Qt.Tool | Qt.WindowTransparentForInput
TIPS = {
    "en": ("Drag a photo from the <b>PomeloO Demo</b> folder and tap <b>{convert}</b> → pick a format",
           "Press <b>{ocr}</b> and drag over the QR code on <b>poster.png</b> → the link is copied",
           "Drag any file to the <b>top of the screen</b> → the notch's drop zones"),
    "id": ("Seret foto dari folder <b>PomeloO Demo</b>, lalu tekan <b>{convert}</b> → pilih format",
           "Tekan <b>{ocr}</b> lalu seret kotak di atas QR di <b>poster.png</b> → link-nya tersalin",
           "Seret file apa saja ke <b>atas layar</b> → zona drop di notch"),
}
TITLE = {"en": ("Try PomeloO", "Free · take it home"), "id": ("Coba PomeloO", "Gratis · bawa pulang")}


def label(pressed: set) -> str | None:
    """What the KeyCaster shows for these held keys, or None for typing (kept private)."""
    mods = [m for m in ORDER if m in pressed]
    keys = sorted(pressed - set(ORDER))
    if keys and not set(mods) - {"Shift"}:  # plain letters, or Shift+letter (capitals)
        return None
    return " + ".join(mods + keys) or None


def pretty(spec: str) -> str:
    """"win+ctrl+t" -> "Win+Ctrl+T"."""
    return "+".join(p.upper() if len(p) == 1 else p.capitalize() for p in spec.split("+"))


def _bottom_center(w: QWidget, lift: int):
    screen = QGuiApplication.screenAt(QCursor.pos()) or QGuiApplication.primaryScreen()
    g = screen.availableGeometry()
    w.adjustSize()
    w.move(g.center().x() - w.width() // 2, g.bottom() - w.height() - lift)


class KeyCaster(QWidget):
    def __init__(self):
        super().__init__(None, FLAGS)
        self.setAttribute(Qt.WA_TranslucentBackground)
        self.setAttribute(Qt.WA_ShowWithoutActivating)
        lay = QVBoxLayout(self)
        lay.setContentsMargins(6, 6, 6, 6)
        self.cap = QLabel()
        self.cap.setStyleSheet("background: rgba(24,24,27,230); border: 1px solid rgba(255,255,255,0.14);"
                               "border-radius: 16px; padding: 10px 28px; font-size: 34px; font-weight: 700;"
                               "color: white;")
        lay.addWidget(self.cap)
        self.held = set()
        self.poll = QTimer(self, interval=30, timeout=self._poll)
        self.hide_timer = QTimer(self, singleShot=True, interval=1200, timeout=self.hide)

    def set_on(self, on: bool):
        self.poll.start() if on else (self.poll.stop(), self.hide())

    def show_keys(self, text: str):
        self.cap.setText(text)
        _bottom_center(self, 24)
        self.show()
        self.raise_()
        self.hide_timer.start()

    def _poll(self):
        down = ctypes.windll.user32.GetAsyncKeyState
        pressed = {name for vk, name in VKS if down(vk) & 0x8000}
        if pressed and not pressed <= self.held and (text := label(pressed)):  # a new key went down
            self.held = pressed
            return self.show_keys(text)
        if pressed and self.isVisible():
            self.hide_timer.start()  # still held: keep it up
        elif not pressed:
            self.held = set()


def _qr_pixmap(text: str, px: int) -> QPixmap:
    import numpy as np
    import zxingcpp
    a = np.ascontiguousarray(np.array(zxingcpp.create_barcode(text, zxingcpp.BarcodeFormat.QRCode).to_image(scale=8)))
    img = QImage(a.data, a.shape[1], a.shape[0], a.strides[0], QImage.Format_Grayscale8).copy()
    return QPixmap.fromImage(img).scaled(px, px, Qt.KeepAspectRatio, Qt.FastTransformation)


class Attract(QWidget):
    def __init__(self, lang: str):
        super().__init__(None, FLAGS)
        self.setAttribute(Qt.WA_TranslucentBackground)
        self.setAttribute(Qt.WA_ShowWithoutActivating)
        outer = QVBoxLayout(self)
        outer.setContentsMargins(0, 0, 0, 0)
        card = QFrame(objectName="card")
        card.setStyleSheet("#card { background: rgba(24,24,27,238); border: 1px solid rgba(255,255,255,0.14);"
                           " border-radius: 22px; } QLabel { color: white; background: transparent; }")
        outer.addWidget(card)
        # a corner card, so whatever is on screen (the website, a folder) stays visible
        row = QHBoxLayout(card)
        row.setContentsMargins(18, 14, 14, 14)
        row.setSpacing(14)
        text = QVBoxLayout()
        title, free = TITLE[lang]
        head = QLabel(title)
        head.setStyleSheet("font-size: 19px; font-weight: 800; color: #ff8a00;")
        self.tip = QLabel(wordWrap=True)
        self.tip.setFixedWidth(270)
        self.tip.setMinimumHeight(66)
        self.tip.setStyleSheet("font-size: 15px;")
        self.dots = QLabel()
        self.dots.setStyleSheet("font-size: 12px; color: rgba(255,255,255,0.5);")
        text.addWidget(head)
        text.addWidget(self.tip, 1)
        text.addWidget(self.dots)
        row.addLayout(text)
        qr = QVBoxLayout()
        code = QLabel()
        code.setPixmap(_qr_pixmap(SITE, 92))
        code.setStyleSheet("background: white; border-radius: 8px; padding: 5px;")
        caption = QLabel(f"{free}<br><span style='color:rgba(255,255,255,0.6)'>{SITE[8:]}</span>")
        caption.setAlignment(Qt.AlignCenter)
        caption.setStyleSheet("font-size: 11px;")
        qr.addWidget(code, 0, Qt.AlignCenter)
        qr.addWidget(caption)
        row.addLayout(qr)
        self.tips, self.i = [], 0
        self.rotate = QTimer(self, interval=6000, timeout=self.next_tip)

    def start(self, tips: list[str]):
        self.tips, self.i = tips, -1
        self.next_tip()
        self.adjustSize()
        g = (QGuiApplication.screenAt(QCursor.pos()) or QGuiApplication.primaryScreen()).availableGeometry()
        self.move(g.right() - self.width() - 16, g.bottom() - self.height() - 16)
        self.show()
        self.rotate.start()

    def stop(self):
        self.rotate.stop()
        self.hide()

    def next_tip(self):
        self.i = (self.i + 1) % len(self.tips)
        self.tip.setText(f"{self.i + 1}.  {self.tips[self.i]}")
        self.dots.setText("  ".join("●" if k == self.i else "○" for k in range(len(self.tips))))


# ---------- demo folder ----------
def demo_dir() -> Path:
    return Path(QStandardPaths.writableLocation(QStandardPaths.DesktopLocation)) / "PomeloO Demo"


def _picture(w, h, top, bottom):
    from PIL import Image, ImageDraw
    from .citrus import LOGO
    img = Image.new("RGB", (w, h))
    d = ImageDraw.Draw(img)
    for y in range(h):  # sky gradient
        t = y / h
        d.line([(0, y), (w, y)], fill=tuple(int(a + (b - a) * t) for a, b in zip(top, bottom)))
    d.ellipse((w * .62, h * .18, w * .82, h * .18 + w * .2), fill="#ffd36e")
    d.polygon([(0, h), (w * .3, h * .55), (w * .55, h * .8), (w * .8, h * .5), (w, h * .75), (w, h)], fill="#3a2350")
    logo = Image.open(LOGO).convert("RGBA").resize((h // 4, h // 4))
    img.paste(logo, (w // 16, h // 12), logo)
    return img


def _poster(path: Path):
    """A QR code printed straight onto a colored background, no white margin (the hard case)."""
    import numpy as np
    import zxingcpp
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (1080, 1350), "#3b1d8f")
    d = ImageDraw.Draw(img)
    d.text((540, 150), "PomeloO", fill="white", anchor="mm", font=ImageFont.truetype("segoeuib.ttf", 110))
    d.text((540, 260), "Scan me · Pindai aku", fill="#ffb020", anchor="mm", font=ImageFont.truetype("segoeui.ttf", 56))
    code = zxingcpp.create_barcode(SITE, zxingcpp.BarcodeFormat.QRCode).to_image(scale=1, add_quiet_zones=False)
    img.paste(Image.fromarray(np.array(code)).convert("RGB").resize((620, 620), Image.NEAREST), (230, 400))
    d.text((540, 1180), SITE[8:], fill="white", anchor="mm", font=ImageFont.truetype("segoeui.ttf", 48))
    img.save(path)


def _receipt():
    """A shop receipt: text for Copy Text (OCR), Redact and Searchable PDF."""
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (900, 1100), "white")
    d = ImageDraw.Draw(img)
    d.text((450, 90), "TOKO POMELO", fill="black", anchor="mm", font=ImageFont.truetype("segoeuib.ttf", 54))
    body = ImageFont.truetype("consola.ttf", 32)
    lines = ["Jl. Jeruk Manis No. 12, Denpasar", "29/10/2026  14:05", "-" * 38,
             "Pomelo   2 x 15.000      30.000", "Mangga   1 x 12.000      12.000", "Teh      3 x  5.000      15.000",
             "-" * 38, "TOTAL                    57.000", "Tunai                   100.000",
             "Kembali                  43.000", "", "Kasir: Ayu   HP 0812-3456-7890", "Terima kasih!"]
    for i, line in enumerate(lines):
        d.text((70, 170 + i * 62), line, fill="black", font=body)
    return img


def _song(path: Path, seconds=6):
    import wave

    import numpy as np
    t = np.arange(int(44100 * seconds)) / 44100
    notes = np.array([523.25, 659.25, 783.99, 1046.5, 783.99, 659.25])[(t * 2).astype(int) % 6]
    pcm = (np.sin(2 * np.pi * notes * t) * np.minimum(1, (1 - (t * 2) % 1) * 4) * 12000).astype("<i2")
    with wave.open(str(path), "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(44100)
        w.writeframes(pcm.tobytes())


def _photos(d: Path):
    import zipfile

    import numpy as np
    import pillow_heif
    from PIL import Image
    from .citrus import LOGO
    pillow_heif.register_heif_opener()
    _picture(1600, 1067, (240, 82, 106), (255, 176, 32)).save(d / "photo.heic", "HEIF", quality=80)
    big = np.asarray(_picture(4000, 2667, (40, 120, 220), (150, 220, 255))).astype(np.int16)
    noise = np.random.default_rng(7).integers(-14, 15, big.shape)  # grain, like a real camera: a big file
    Image.fromarray(np.clip(big + noise, 0, 255).astype(np.uint8)).save(d / "holiday.jpg", quality=95)
    for i, (top, bottom, when) in enumerate([((34, 139, 87), (180, 230, 160), "2026:08:17 09:12:00"),
                                             ((90, 40, 160), (255, 140, 120), "2026:08:17 18:40:00"),
                                             ((20, 60, 120), (90, 200, 220), "2026:08:18 07:05:00")], 1):
        img = _picture(1600, 1067, top, bottom)
        exif = img.getexif()
        exif[306] = when
        exif.get_ifd(0x8769)[36867] = when  # DateTimeOriginal: Rename can number them by date taken
        img.save(d / f"trip-{i}.jpg", quality=90, exif=exif)
    _picture(1600, 1067, (255, 120, 80), (255, 220, 150)).save(d / "beach.webp", quality=85)
    sticker = Image.new("RGBA", (900, 900))
    logo = Image.open(LOGO).convert("RGBA").resize((760, 760))
    sticker.paste(logo, (70, 70), logo)
    sticker.save(d / "sticker.png")  # transparent: Add Background
    _receipt().save(d / "receipt.png")
    _poster(d / "poster.png")
    with zipfile.ZipFile(d.parent / "Extras" / "trip-photos.zip", "w") as z:
        for i in (1, 2, 3):
            z.write(d / f"trip-{i}.jpg", f"trip-{i}.jpg")


def _media(d: Path):
    from ..engines.media import ff
    _song(d / "song.wav")
    ff("-i", d / "song.wav", "-b:a", "192k", d / "melody.mp3")
    h264 = ("-c:v", "libx264", "-preset", "veryfast", "-pix_fmt", "yuv420p")
    ff("-f", "lavfi", "-i", "testsrc2=size=1280x720:rate=30:duration=6", "-i", d / "song.wav", "-shortest",
       *h264, "-c:a", "aac", d / "video.mp4")
    (d / "video.srt").write_text("1\n00:00:00,300 --> 00:00:02,800\nHalo dari PomeloO!\n\n"
                                 "2\n00:00:03,000 --> 00:00:05,800\nSubtitle ini dibakar ke videonya.\n", encoding="utf-8")
    ff("-f", "lavfi", "-i", "gradients=size=1280x720:rate=30:speed=0.03:nb_colors=3:c0=0xf0526a:c1=0xffb020"
       ":c2=0x3b1d8f", "-t", "5", *h264, d / "clip-2.mp4")


def _documents(d: Path):
    import pymupdf as fitz
    from openpyxl import Workbook
    from pptx import Presentation
    holiday = str(d.parent / "Photos" / "holiday.jpg")
    doc = fitz.open()
    for line in ["PomeloO demo report", "Page two: split me, number me, watermark me"]:
        page = doc.new_page()
        page.insert_text((72, 96), line, fontsize=28)
        page.insert_text((72, 140), "Everything happens on this PC. Nothing is uploaded.", fontsize=14)
        page.insert_image(fitz.Rect(72, 180, 523, 480), filename=holiday)
    doc.save(d / "report.pdf")
    doc = fitz.open()
    doc.new_page().insert_text((72, 96), "Chapter 2: merge me with report.pdf", fontsize=26)
    doc.save(d / "chapter-2.pdf")
    doc = fitz.open()  # a scan: only a picture, no text (Searchable PDF fixes that)
    doc.new_page(width=450, height=550).insert_image(fitz.Rect(0, 0, 450, 550),
                                                     filename=str(d.parent / "Photos" / "receipt.png"))
    doc.save(d / "scan.pdf")
    (d / "notes.md").write_text("# Shopping list\n\n- Pomelo\n- Mango\n- *Lots* of tea\n\n"
                                "| Day | Plan |\n|---|---|\n| Sat | Exhibition |\n", encoding="utf-8")
    prs = Presentation()
    for title, points in [("PomeloO", ["Drag a file, tap a key, pick a format"]),
                          ("What it does", ["275+ conversions", "34 tools", "8 actions"]),
                          ("Private", ["Everything happens on this PC", "Nothing is uploaded"])]:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = title
        tf = s.placeholders[1].text_frame
        tf.text = points[0]
        for p in points[1:]:
            tf.add_paragraph().text = p
    prs.save(d / "slides.pptx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Item", "Qty", "Price", "Total"])
    for r, (item, qty, price) in enumerate([("Pomelo", 2, 15000), ("Mango", 1, 12000), ("Tea", 3, 5000)], 2):
        ws.append([item, qty, price, f"=B{r}*C{r}"])
    ws.append(["Total", None, None, "=SUM(D2:D4)"])
    wb.save(d / "budget.xlsx")


DEMO = {"Photos": _photos, "Videos & Audio": _media, "Documents": _documents}


def make_demo(d: Path) -> int:
    """Empty `d` (only if it is a demo folder) and fill it with sample files. Returns files that couldn't
    be removed (open in another app). Never touches a folder without the MARK file."""
    import shutil
    if d.exists() and not (d / MARK).exists():
        raise RuntimeError(f"{d} already exists and isn't a PomeloO demo folder, so it was left alone.")
    stuck = 0
    if d.exists():
        for p in d.iterdir():
            if p.name == MARK:
                continue
            try:
                shutil.rmtree(p) if p.is_dir() else p.unlink()
            except OSError:
                stuck += 1
    d.mkdir(parents=True, exist_ok=True)
    mark = d / MARK
    if not mark.exists():
        mark.write_text("Made by PomeloO's Exhibition mode. Tray -> Reset demo folder empties this folder.\n")
        ctypes.windll.kernel32.SetFileAttributesW(str(mark), 2)  # hidden
    for sub in (*DEMO, "Extras"):
        (d / sub).mkdir(exist_ok=True)
    for sub, make in DEMO.items():  # Photos first: the others reuse its pictures
        make(d / sub)
    return stuck


class Expo(QObject):
    """Needs from the app: settings, notify(text), busy(expo=False), screen_text."""
    made = Signal(str, int, str)  # folder, stuck files, error

    def __init__(self, app):
        super().__init__()
        self.app = app
        self.lang = "id" if QLocale.system().language() == QLocale.Indonesian else "en"
        self.caster, self.attract = KeyCaster(), Attract(self.lang)
        self.idle = QTimer(self, interval=1000, timeout=self._tick)
        self.made.connect(self._made)
        self.set_on(app.settings.get("exhibition", False))

    def set_on(self, on: bool):
        self.caster.set_on(on)
        self.idle.start() if on else (self.idle.stop(), self.attract.stop())
        self.app.screen_text.toast.big = on

    def _tick(self, idle=None):
        from ..update import idle_seconds
        idle = idle_seconds() if idle is None else idle
        if idle >= IDLE_S and not self.attract.isVisible() and not self.app.busy(expo=False):
            keys = {v: k for k, v in self.app.settings["keys"].items()}
            tips = [t.format(convert=pretty(keys.get("convert", "shift")), ocr=pretty(self.app.settings["ocr_hotkey"]))
                    for t in TIPS[self.lang]]
            self.attract.start(tips)
        elif idle < 1 and self.attract.isVisible():
            self.attract.stop()

    def reset_folder(self):
        d = demo_dir()

        def work():
            try:
                self.made.emit(str(d), make_demo(d), "")
            except Exception as e:  # shown to the user, never swallowed
                self.made.emit(str(d), 0, str(e) or e.__class__.__name__)
        threading.Thread(target=work, daemon=True).start()

    def _made(self, folder, stuck, err):
        if err:
            return self.app.notify(f"Demo folder: {err}")
        os.startfile(folder)
        self.app.notify("Demo folder is ready." + (f" {stuck} file(s) were open elsewhere and stayed." if stuck else ""))


if __name__ == "__main__":
    import sys
    import tempfile

    assert label({"Win", "Ctrl", "T"}) == "Win + Ctrl + T"
    assert label({"Ctrl", "Shift", "Space"}) == "Ctrl + Shift + Space"
    assert label({"Shift"}) == "Shift" and label({"Ctrl", "V"}) == "Ctrl + V"
    assert label({"T"}) is None and label({"Shift", "T"}) is None and label(set()) is None
    assert pretty("win+ctrl+t") == "Win+Ctrl+T"
    from PySide6.QtWidgets import QApplication
    qapp = QApplication(sys.argv)
    from PIL import Image
    from .snip import read_image
    with tempfile.TemporaryDirectory() as tmp:
        d = Path(tmp) / "PomeloO Demo"
        assert make_demo(d) == 0 and make_demo(d) == 0  # the second run is a reset
        names = sorted(p.relative_to(d).as_posix() for p in d.rglob("*") if p.is_file())
        print(len(names), "files:", ", ".join(names))
        assert len(names) == 22 and "Videos & Audio/video.srt" in names and "Extras/trip-photos.zip" in names, names
        assert read_image(Image.open(d / "Photos" / "poster.png").convert("RGB")) == (SITE, True)
        assert "TOKO POMELO" in read_image(Image.open(d / "Photos" / "receipt.png").convert("RGB"))[0]
        from ..rename import taken
        assert taken(d / "Photos" / "trip-2.jpg").hour == 18
        from ..catalog import targets, tools
        for rel in names:  # every sample can be converted or has tools in PomeloO
            ext = Path(rel).suffix.lstrip(".")
            assert rel == MARK or rel.endswith(".srt") or tools(ext) or targets(ext), rel
        other = Path(tmp) / "Mine"
        other.mkdir()
        (other / "keep.txt").write_text("mine")
        try:
            make_demo(other)
            raise AssertionError("emptied a folder that isn't a demo folder")
        except RuntimeError:
            assert (other / "keep.txt").exists()
    print("expo self-check OK")
