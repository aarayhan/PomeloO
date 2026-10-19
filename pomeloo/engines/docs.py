"""PDFs, rich-text documents, slides and spreadsheets."""
import html
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

import pymupdf as fitz
import pypandoc

from ..catalog import stem_of, unique

_NOWIN = getattr(subprocess, "CREATE_NO_WINDOW", 0)
PANDOC_IN = {"docx": "docx", "md": "markdown", "rtf": "rtf", "html": "html", "epub": "epub"}
PANDOC_OUT = {"docx": "docx", "md": "gfm", "html": "html", "epub": "epub3", "txt": "plain"}
FONTS = Path(r"C:\Windows\Fonts")
# Segoe UI covers bullets, dashes, Latin-extended text; the PDF base-14 fonts do not.
CSS = ("@font-face{font-family:ui;src:url(segoeui.ttf)} @font-face{font-family:ui;font-weight:bold;src:url(segoeuib.ttf)}"
       " @font-face{font-family:ui;font-style:italic;src:url(segoeuii.ttf)}"
       " @font-face{font-family:mono;src:url(consola.ttf)}"
       " body{font-family:ui,sans-serif;font-size:11pt;line-height:1.4} table{border-collapse:collapse}"
       " td,th{border:1px solid #999;padding:2px 4px;font-size:9pt} pre,code{font-family:mono,monospace}")


def pandoc(*args, cwd=None):
    p = subprocess.run([pypandoc.get_pandoc_path(), *map(str, args)], capture_output=True,
                       text=True, encoding="utf-8", errors="replace", cwd=cwd, creationflags=_NOWIN)
    if p.returncode:
        raise RuntimeError("pandoc: " + p.stderr.strip()[-400:])


def _pandoc_to(src: Path, fmt_in: str, target: str, dst: Path, title: str):
    extra = []
    if target == "epub":  # e-readers show this as the book title
        extra = ["-s", "--metadata", f"title={title}"]
    if target == "html":  # pagetitle = browser tab only, no extra heading in the page
        extra = ["-s", "--metadata", f"pagetitle={title}", "--embed-resources"]
    pandoc(src, "-f", fmt_in, "-t", PANDOC_OUT[target], "-o", dst,
           "--resource-path", src.parent, *extra)


def html_to_pdf(html_text: str, dst: Path, base: Path):
    """Lay out HTML onto A4 pages with PyMuPDF's Story engine (no LibreOffice needed)."""
    archive = fitz.Archive(str(base))
    if FONTS.exists():
        archive.add(str(FONTS))
    story = fitz.Story(html=html_text, user_css=CSS, archive=archive)
    writer = fitz.DocumentWriter(str(dst))
    rect = fitz.paper_rect("a4")
    more = True
    while more:
        dev = writer.begin_page(rect)
        more, _ = story.place(rect + (50, 50, -50, -50))
        story.draw(dev)
        writer.end_page()
    writer.close()


def render_pages(pdf: Path, dst: Path, ext: str, dpi: int = 300) -> Path:
    """One page -> dst; several -> folder `<stem>/` with `<stem> N.ext`."""
    doc = fitz.open(pdf)
    if len(doc) == 1:
        doc[0].get_pixmap(dpi=dpi).save(dst, jpg_quality=92)
        return dst
    folder = unique(dst.with_suffix(""))
    folder.mkdir()
    for i, page in enumerate(doc, 1):
        page.get_pixmap(dpi=dpi).save(folder / f"{folder.name} {i}.{ext}", jpg_quality=92)
    return folder


# ---------- PDF ----------
def _pdf_to_pptx(src: Path, dst: Path, tmp: Path):
    from pptx import Presentation
    from pptx.util import Emu
    doc = fitz.open(src)
    prs = Presentation()
    w, h = doc[0].rect.width, doc[0].rect.height
    prs.slide_width, prs.slide_height = Emu(int(w * 12700)), Emu(int(h * 12700))
    for i, page in enumerate(doc):
        img = tmp / f"p{i}.png"
        page.get_pixmap(dpi=200).save(img)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(img), 0, 0, prs.slide_width, prs.slide_height)
    prs.save(dst)


def _pdf_to_xlsx(src: Path, dst: Path):
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    doc = fitz.open(src)
    for pno, page in enumerate(doc, 1):
        for tno, table in enumerate(page.find_tables().tables, 1):
            ws = wb.create_sheet(f"Page {pno} Table {tno}"[:31])
            for row in table.extract():
                ws.append([c if c is not None else "" for c in row])
    if not wb.worksheets:  # ponytail: no ruled tables found -> one text line per row
        ws = wb.create_sheet("Text")
        for page in doc:
            for line in page.get_text().splitlines():
                ws.append([line])
    wb.save(dst)


def pdf_convert(src: Path, dst: Path, target: str, title: str) -> Path:
    with tempfile.TemporaryDirectory() as t:
        tmp = Path(t)
        if target == "docx":
            from pdf2docx import Converter
            cv = Converter(str(src))
            try:
                cv.convert(str(dst))
            finally:
                cv.close()
        elif target == "pptx":
            _pdf_to_pptx(src, dst, tmp)
        elif target == "xlsx":
            _pdf_to_xlsx(src, dst)
        elif target in ("jpg", "png"):
            return render_pages(src, dst, target)
        elif target == "pdf":
            shutil.copyfile(src, dst)
        elif target == "txt":
            dst.write_text("\n\n".join(p.get_text() for p in fitz.open(src)), encoding="utf-8")
        else:  # md, epub, html go through Markdown
            import pymupdf4llm
            md = pymupdf4llm.to_markdown(str(src))
            if target == "md":
                dst.write_text(md, encoding="utf-8")
            else:
                (tmp / "doc.md").write_text(md, encoding="utf-8")
                _pandoc_to(tmp / "doc.md", "markdown", target, dst, title)
    return dst


# ---------- Pages / RTFD / slides ----------
def pages_preview(src: Path, tmp: Path) -> Path:
    """`.pages` is a zip; Pages embeds a preview PDF (or JPEG) we can read."""
    with zipfile.ZipFile(src) as z:
        names = z.namelist()
        pdf = next((n for n in names if n.lower().endswith("preview.pdf")), None)
        if pdf:
            out = tmp / "pages.pdf"
            out.write_bytes(z.read(pdf))
            return out
        jpg = next((n for n in names if n.lower().startswith("preview") and n.lower().endswith(".jpg")), None)
        if not jpg:
            raise RuntimeError("This .pages file has no embedded preview to convert.")
        (tmp / "p.jpg").write_bytes(z.read(jpg))
        out = tmp / "pages.pdf"
        out.write_bytes(fitz.open(tmp / "p.jpg").convert_to_pdf())
        return out


def _soffice_to_pdf(exe, src: Path, tmp: Path) -> Path:
    profile = (tmp / "lo").as_uri()  # throwaway profile: never touches a running LibreOffice
    subprocess.run([str(exe), f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf",
                    "--outdir", str(tmp), str(src)], capture_output=True, timeout=600, creationflags=_NOWIN)
    out = tmp / f"{src.stem}.pdf"
    if not out.exists():
        raise RuntimeError("LibreOffice could not convert this presentation.")
    return out


def slides_to_pdf(src: Path, tmp: Path) -> Path:
    """Full-fidelity rendering by a real Office engine: installed LibreOffice, WPS, PowerPoint,
    else PomeloO's private LibreOffice (downloaded once, see engines/libreoffice.py)."""
    for exe in (shutil.which("soffice"), r"C:\Program Files\LibreOffice\program\soffice.exe",
                r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if exe and Path(exe).exists():
            return _soffice_to_pdf(exe, src, tmp)
    out = tmp / f"{src.stem}.pdf"
    for progid in ("KWPP.Application", "PowerPoint.Application"):  # WPS Presentation, then MS PowerPoint
        if _office_to_pdf(progid, src, out):
            return out
    from . import libreoffice
    return _soffice_to_pdf(libreoffice.ensure(), src, tmp)


def _office_to_pdf(progid: str, src: Path, out: Path) -> bool:
    """Export through an installed Office suite over COM (same object model in WPS and PowerPoint).
    Only the presentation we opened is closed; the app quits only if nothing else of the user's is open."""
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()  # jobs run on worker threads
    try:
        app = win32com.client.DispatchEx(progid)
    except Exception:
        pythoncom.CoUninitialize()
        return False
    try:
        pres = app.Presentations.Open(str(src), True, False, False)  # ReadOnly, Untitled, WithWindow
        try:
            pres.SaveAs(str(out), 32)  # ppSaveAsPDF
        finally:
            pres.Close()
        return out.exists()
    except Exception:
        return False
    finally:
        try:
            if app.Presentations.Count == 0 and not app.Visible:
                app.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()


# ---------- Spreadsheets ----------
_NUM = re.compile(r"-?(?:0|[1-9]\d*)(?:\.\d+)?")


def _number(v):
    """'3' -> 3, '1.2' -> 1.2 so spreadsheets can calculate; '007', '+62..' and text stay as typed."""
    if isinstance(v, str) and _NUM.fullmatch(v.strip()):
        return float(v) if "." in v else int(v)
    return v


def read_sheets(src: Path, ext: str) -> dict:
    import pandas as pd
    if ext in ("csv", "tsv"):
        df = pd.read_csv(src, sep="\t" if ext == "tsv" else ",", header=None, dtype=str,
                         keep_default_na=False, encoding_errors="replace")
        return {stem_of(src)[:31] or "Sheet1": df.map(_number)}
    return pd.read_excel(src, sheet_name=None, header=None,
                         engine={"xls": "xlrd", "ods": "odf"}.get(ext, "openpyxl"))


def sheet_convert(src: Path, dst: Path, ext: str, target: str):
    import pandas as pd
    sheets = read_sheets(src, ext)
    if target in ("xlsx", "ods"):
        with pd.ExcelWriter(dst, engine="odf" if target == "ods" else "openpyxl") as w:
            for name, df in sheets.items():
                df.to_excel(w, sheet_name=str(name)[:31], index=False, header=False)
    elif target in ("csv", "tsv"):  # ponytail: first sheet only, like Numbers' CSV export
        next(iter(sheets.values())).to_csv(dst, sep="\t" if target == "tsv" else ",",
                                           index=False, header=False)
    else:  # pdf
        body = "".join(f"<h3>{html.escape(str(n))}</h3>"
                       + df.to_html(index=False, header=False, na_rep="") for n, df in sheets.items())
        html_to_pdf(f"<html><body>{body}</body></html>", dst, src.parent)


# ---------- entry ----------
def _txt_as_html(src: Path) -> str:
    text = src.read_text(encoding="utf-8", errors="replace")
    paras = (html.escape(p).replace("\n", "<br/>") for p in text.split("\n\n"))
    return "<html><body>" + "".join(f"<p>{p}</p>" for p in paras) + "</body></html>"


def convert(src: Path, dst: Path, ext: str, target: str) -> Path:
    title = stem_of(src)
    with tempfile.TemporaryDirectory() as t:
        tmp = Path(t)
        if ext == "pdf":
            return pdf_convert(src, dst, target, title)
        if ext == "pages":
            return pdf_convert(pages_preview(src, tmp), dst, target, title)
        if ext in ("ppt", "pptx"):
            pdf = slides_to_pdf(src, tmp)
            if target == "pdf":
                shutil.copyfile(pdf, dst)
                return dst
            return render_pages(pdf, dst, target)
        if ext in ("xlsx", "xls", "ods", "csv", "tsv"):
            sheet_convert(src, dst, ext, target)
            return dst
        # Rich text: docx md rtf rtfd txt html epub
        if ext == "rtfd":
            src, ext = src / "TXT.rtf", "rtf"
        if ext == "txt":
            (tmp / "doc.html").write_text(_txt_as_html(src), encoding="utf-8")
            src, ext = tmp / "doc.html", "html"
        if target in PANDOC_OUT:
            _pandoc_to(src, PANDOC_IN[ext], target, dst, title)
            return dst
        # pdf / png / jpg: pandoc -> HTML fragment (media extracted into tmp) -> Story PDF.
        # No "-s": pandoc's default CSS sets table{display:block}, which stacks every cell.
        pandoc(src.resolve(), "-f", PANDOC_IN[ext], "-t", "html", "--extract-media=.",
               "--resource-path", src.parent.resolve(), "-o", "doc.html", cwd=tmp)
        pdf = dst if target == "pdf" else tmp / "out.pdf"
        body = (tmp / "doc.html").read_text(encoding="utf-8")
        html_to_pdf(f"<html><body>{body}</body></html>", pdf, tmp)
        return dst if target == "pdf" else render_pages(pdf, dst, target)
